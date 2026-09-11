"""One-time bulk import of existing inventory rows from a CSV or XLSX export.

Usage:
    python import_data.py path/to/inventory_export.csv
    python import_data.py path/to/inventory_export.xlsx
    python import_data.py path/to/inventory_export.csv --dry-run

Column headers are matched case-insensitively and ignoring punctuation, so
both the plain schema names (item_id, storage_location, ...) and the more
readable spreadsheet-style headers (Item ID, Storage Location, "Photo On
File?", ...) are recognized. See FIELD_ALIASES below for the exact list.

- `item_id` is preserved from the source file when present (zero-padded to
  4 digits if it's numeric). Rows with no item_id get the next sequential id.
- Rows missing `name` are skipped and reported.
- Rows with an `availability_status` outside the allowed set are imported
  with status defaulted to "In Stock" and flagged in the summary so you can
  fix them by hand afterward.
- `photo_on_file` / `listed` accept common truthy strings (1, true, yes, y).
- Numeric fields (costs/prices/profit) that can't be parsed as a number are
  imported as blank and flagged in the summary.
"""

import argparse
import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from app import db  # noqa: E402

TRUTHY = {"1", "true", "yes", "y", "t"}

# canonical field -> list of header spellings that should map to it (each
# entry is compared after _clean() normalizes both sides)
FIELD_ALIASES = {
    "item_id": ["item_id", "item id", "id"],
    "name": ["name", "item_name", "item name description", "item name", "description"],
    "category": ["category"],
    "storage_location": ["storage_location", "storage location"],
    "availability_status": ["availability_status", "status"],
    "condition": ["condition"],
    "photo_on_file": ["photo_on_file", "photo on file"],
    "notes": ["notes"],
    "acquired_date": ["acquired_date", "acquired date"],
    "acquired_from": ["acquired_from", "acquired from source", "acquired from"],
    "acquisition_cost": ["acquisition_cost", "acquisition cost"],
    "listed": ["listed"],
    "listing_platform": ["listing_platform", "listing platform"],
    "listing_price": ["listing_price", "listing price"],
    "listing_date": ["listing_date", "listing date"],
    "sold_date": ["sold_date", "sold date"],
    "sold_price": ["sold_price", "sold price"],
    "fees_shipping_cost": ["fees_shipping_cost", "fees shipping cost"],
    "net_profit": ["net_profit", "net profit"],
    "final_disposition": ["final_disposition", "final disposition"],
}

NUMERIC_FIELDS = {"acquisition_cost", "listing_price", "sold_price", "fees_shipping_cost", "net_profit"}
BOOL_FIELDS = {"photo_on_file", "listed"}


def _clean(header: str) -> str:
    header = (header or "").strip().lower()
    header = header.replace("?", "")
    header = re.sub(r"[()]", "", header)
    header = re.sub(r"[\s/_]+", " ", header).strip()
    return header


_ALIAS_LOOKUP = {
    _clean(alias): field
    for field, aliases in FIELD_ALIASES.items()
    for alias in aliases
}


def _stringify(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value).strip()


def normalize_row(raw: dict) -> dict:
    mapped = {}
    for header, value in raw.items():
        field = _ALIAS_LOOKUP.get(_clean(header))
        if field:
            mapped[field] = _stringify(value)

    row = {field: mapped.get(field, "") for field in FIELD_ALIASES}

    for field in BOOL_FIELDS:
        row[field] = row[field].strip().lower() in TRUTHY

    return row


def read_csv(path: Path):
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for raw in reader:
            yield normalize_row(raw)


def read_xlsx(path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    for values in rows:
        if values is None or all(v is None for v in values):
            continue
        raw = dict(zip(header, values))
        yield normalize_row(raw)


def _parse_numeric_fields(row: dict, warnings: list, item_label: str) -> dict:
    parsed = dict(row)
    for field in NUMERIC_FIELDS:
        raw = row[field]
        if not raw:
            parsed[field] = None
            continue
        try:
            parsed[field] = float(raw)
        except ValueError:
            warnings.append((item_label, field, raw))
            parsed[field] = None
    return parsed


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", help="Path to the CSV or XLSX file to import")
    parser.add_argument("--dry-run", action="store_true", help="Parse and validate without writing to the database")
    args = parser.parse_args()

    path = Path(args.source)
    if not path.exists():
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)

    if path.suffix.lower() == ".csv":
        rows = list(read_csv(path))
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        rows = list(read_xlsx(path))
    else:
        print(f"Unsupported file type: {path.suffix}", file=sys.stderr)
        sys.exit(1)

    db.init_db()
    conn = db.get_connection()

    imported = 0
    skipped_missing_name = 0
    skipped_duplicate_id = []
    status_warnings = []
    numeric_warnings = []

    try:
        for row in rows:
            if not row["name"]:
                skipped_missing_name += 1
                continue

            item_id = row["item_id"]
            if item_id and item_id.isdigit():
                item_id = item_id.zfill(4)

            if item_id and db.item_id_exists(conn, item_id):
                skipped_duplicate_id.append(item_id)
                continue

            label = item_id or f"(auto, {row['name']!r})"

            status = row["availability_status"]
            if status not in db.ALLOWED_STATUSES:
                if status:
                    status_warnings.append((label, status))
                status = "In Stock"
            row["availability_status"] = status

            row = _parse_numeric_fields(row, numeric_warnings, label)

            if args.dry_run:
                imported += 1
                continue

            if item_id:
                row["item_id"] = item_id
                extra_cols = [c for c, _ in db.EXTRA_COLUMNS]
                conn.execute(
                    f"""
                    INSERT INTO items
                        (item_id, name, category, storage_location, availability_status,
                         condition, photo_on_file, notes, {', '.join(extra_cols)})
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, {', '.join(['?'] * len(extra_cols))})
                    """,
                    (
                        item_id,
                        row["name"],
                        row["category"] or None,
                        row["storage_location"] or None,
                        status,
                        row["condition"] or None,
                        1 if row["photo_on_file"] else 0,
                        row["notes"] or None,
                        *[
                            (1 if row[c] else 0) if c in BOOL_FIELDS else (row[c] or None)
                            for c, _ in db.EXTRA_COLUMNS
                        ],
                    ),
                )
            else:
                db.create_item(conn, row)
            imported += 1

        if not args.dry_run:
            conn.commit()
    finally:
        conn.close()

    print(f"Imported: {imported}")
    print(f"Skipped (missing name): {skipped_missing_name}")
    if skipped_duplicate_id:
        print(f"Skipped (item_id already exists): {len(skipped_duplicate_id)} -> {skipped_duplicate_id}")
    if status_warnings:
        print(f"Rows with invalid availability_status defaulted to 'In Stock': {len(status_warnings)}")
        for item_id, bad_status in status_warnings:
            print(f"  - {item_id}: {bad_status!r}")
    if numeric_warnings:
        print(f"Rows with unparseable numeric values (left blank): {len(numeric_warnings)}")
        for label, field, raw in numeric_warnings:
            print(f"  - {label}: {field} = {raw!r}")
    if args.dry_run:
        print("(dry run — no rows written)")


if __name__ == "__main__":
    main()
